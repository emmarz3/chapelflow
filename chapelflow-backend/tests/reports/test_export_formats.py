import pytest


@pytest.mark.django_db
class TestReportExportEngine:
    """
    Spec section 16: "Do not claim Excel/PDF support if the task actually
    only generates CSV." These exercise the real per-format serializers in
    apps.reports.services, not just that a ReportJob row was created.
    """

    def test_export_csv_produces_valid_csv_bytes(self):
        from apps.reports.services import export_rows_csv

        rows = [{"status": "ACTIVE", "count": 5}, {"status": "INACTIVE", "count": 2}]
        content = export_rows_csv(rows)
        text = content.decode("utf-8")
        assert "status,count" in text
        assert "ACTIVE,5" in text

    def test_export_csv_handles_empty_rows(self):
        from apps.reports.services import export_rows_csv

        assert export_rows_csv([]) == b""

    def test_export_excel_produces_a_real_readable_workbook(self):
        import io

        from openpyxl import load_workbook

        from apps.reports.services import export_rows_excel

        rows = [{"status": "ACTIVE", "count": 5}, {"status": "INACTIVE", "count": 2}]
        content = export_rows_excel(rows)

        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        assert header == ["status", "count"]
        data_row = [cell.value for cell in ws[2]]
        assert data_row == ["ACTIVE", 5]

    def test_export_excel_handles_empty_rows(self):
        from openpyxl import load_workbook
        import io

        from apps.reports.services import export_rows_excel

        content = export_rows_excel([])
        wb = load_workbook(io.BytesIO(content))
        assert wb.active.max_row == 1  # no header/data written, but a valid empty workbook

    def test_export_pdf_produces_a_real_pdf(self):
        from apps.reports.services import export_rows_pdf

        rows = [{"status": "ACTIVE", "count": 5}]
        content = export_rows_pdf(rows, title="Membership Report")
        assert content.startswith(b"%PDF")  # genuine PDF file signature, not a stub

    def test_export_pdf_handles_empty_rows_without_crashing(self):
        from apps.reports.services import export_rows_pdf

        content = export_rows_pdf([], title="Empty Report")
        assert content.startswith(b"%PDF")

    def test_export_report_rows_dispatches_correctly_per_format(self):
        from apps.reports.services import export_report_rows

        rows = [{"a": 1}]
        csv_bytes, csv_ct, csv_ext = export_report_rows("CSV", rows)
        assert csv_ct == "text/csv" and csv_ext == "csv"

        xlsx_bytes, xlsx_ct, xlsx_ext = export_report_rows("EXCEL", rows)
        assert xlsx_ext == "xlsx"
        assert xlsx_ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        pdf_bytes, pdf_ct, pdf_ext = export_report_rows("PDF", rows, title="X")
        assert pdf_ct == "application/pdf" and pdf_ext == "pdf"
        assert pdf_bytes.startswith(b"%PDF")

    def test_export_report_rows_rejects_unknown_format(self):
        from apps.reports.services import export_report_rows

        with pytest.raises(ValueError):
            export_report_rows("WORD_DOC", [{"a": 1}])


@pytest.mark.django_db
class TestReportJobEndToEndPerFormat:
    """
    Full loop through ReportJob -> run_report_job -> real file written via
    StorageService, for each format — catches the exact bug this session
    fixed (job.export_format being ignored and CSV always written).
    """

    @pytest.fixture
    def _membership_data(self, branch_a):
        from apps.members.models import Member

        Member.objects.create(branch=branch_a, first_name="A", last_name="One", membership_status="ACTIVE")
        Member.objects.create(branch=branch_a, first_name="B", last_name="Two", membership_status="ACTIVE")

    @pytest.mark.parametrize("export_format,expected_ext", [("CSV", "csv"), ("EXCEL", "xlsx"), ("PDF", "pdf")])
    def test_run_report_job_writes_the_requested_format(self, branch_a, chapel_admin_a, _membership_data, export_format, expected_ext):
        from apps.reports.models import ReportJob, ReportJobStatus
        from apps.reports.tasks import run_report_job

        job = ReportJob.objects.create(
            branch=branch_a, requested_by=chapel_admin_a, report_type="MEMBERSHIP",
            export_format=export_format, filters={},
        )
        run_report_job(str(job.id))  # CELERY_TASK_ALWAYS_EAGER in test settings, but call directly to be explicit

        job.refresh_from_db()
        assert job.status == ReportJobStatus.COMPLETE
        assert job.file_url
        assert job.file_url.endswith(f".{expected_ext}")

    def test_run_report_job_marks_failed_on_unknown_report_type(self, branch_a, chapel_admin_a):
        from apps.reports.models import ReportJob, ReportJobStatus
        from apps.reports.tasks import run_report_job

        job = ReportJob.objects.create(
            branch=branch_a, requested_by=chapel_admin_a, report_type="VOLUNTEERS",  # no generator registered
            export_format="CSV", filters={},
        )
        with pytest.raises(ValueError):
            run_report_job(str(job.id))

        job.refresh_from_db()
        assert job.status == ReportJobStatus.FAILED
        assert job.error_message


@pytest.mark.django_db
class TestReportsAPIFormats:
    def test_formats_endpoint_lists_all_three_implemented_formats(self, api_client, chapel_admin_a, seed_member_permissions):
        api_client.force_authenticate(user=chapel_admin_a)
        response = api_client.get("/api/v1/reports/jobs/formats/")
        assert response.status_code == 200
        assert set(response.data["data"].keys()) == {"CSV", "EXCEL", "PDF"}

    def test_creating_a_report_job_with_excel_format_is_accepted(self, api_client, chapel_admin_a, branch_a, seed_member_permissions):
        api_client.force_authenticate(user=chapel_admin_a)
        response = api_client.post("/api/v1/reports/jobs/", {
            "branch": str(branch_a.id), "report_type": "MEMBERSHIP", "export_format": "EXCEL",
        })
        assert response.status_code == 201

        from apps.reports.models import ReportJob
        job = ReportJob.objects.get(id=response.data["data"]["id"])
        # CELERY_TASK_ALWAYS_EAGER in test settings runs this synchronously.
        assert job.file_url.endswith(".xlsx")
